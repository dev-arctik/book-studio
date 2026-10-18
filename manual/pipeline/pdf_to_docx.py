# manual/pipeline/pdf_to_docx.py
# Turns a finished, text-based book PDF (not a scan) into an editable Word file next to it: books/final/<book>.pdf → <book>.docx.
# The text itself is rebuilt by pdf_text.py (paragraphs, headings, bold/italic, verse breaks); this file only writes Word.
# Page numbers are dropped (Word numbers its own pages).
# Usage: poetry run python pipeline/pdf_to_docx.py charu-vilas-vol2_en charu-vilas-vol2_hi

import sys; sys.dont_write_bytecode = True  # keep pipeline/ free of __pycache__ clutter
import local_env

from collections import Counter

import pymupdf
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from pdf_text import DEVANAGARI_RE, read_pdf


def set_cs_size(rpr, size: float) -> None:
    """Font size for complex-script text (Devanagari) — Word keeps it apart from the Latin size (w:sz)."""
    el = rpr.find(qn("w:szCs"))
    if el is None:
        el = OxmlElement("w:szCs")
        rpr.append(el)
    el.set(qn("w:val"), str(round(size * 2)))


def set_fonts(doc: Document, latin: str, devanagari: str, size: float) -> None:
    """Default font for the whole document: Word picks the complex-script font (w:cs) for Devanagari."""
    rpr = doc.styles["Normal"].element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for attr, name in (("w:ascii", latin), ("w:hAnsi", latin), ("w:cs", devanagari)):
        fonts.set(qn(attr), name)
    doc.styles["Normal"].font.size = Pt(size)
    set_cs_size(rpr, size)
    for style in ("Title", "Heading 1", "Heading 2"):
        doc.styles[style].element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), devanagari)


def add_run(paragraph, r: dict, body: float, kind: str) -> None:
    if r.get("break"):
        paragraph.add_run().add_break(WD_BREAK.LINE)
        return
    run = paragraph.add_run(r["text"])
    run.bold = r["bold"] or None
    run.italic = r["italic"] or None
    if kind == "p" and r["size"] != body:
        run.font.size = Pt(r["size"])  # dates and small notes keep their smaller size
        set_cs_size(run.element.get_or_add_rPr(), r["size"])


def convert(book: str) -> None:
    pdf_path = local_env.find_book(book)
    pdf = pymupdf.open(pdf_path)
    paras, body = read_pdf(pdf)
    hindi = any(DEVANAGARI_RE.search(p["text"]) for p in paras[:50])

    doc = Document()
    set_fonts(doc, "Arial", "Noto Serif Devanagari", body)
    doc.core_properties.title = pdf.metadata.get("title") or book
    for i, p in enumerate(paras):
        # A part heading (Heading 1) opens a new page, as in the PDF
        if p["kind"] == "Heading 1" and i > 0 and paras[i - 1]["kind"] != "Heading 1":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        para = doc.add_paragraph(style=None if p["kind"] == "p" else p["kind"])
        for r in p["runs"]:
            add_run(para, r, body, p["kind"])
    out = pdf_path.with_suffix(".docx")
    doc.save(out)
    kinds = Counter(p["kind"] for p in paras)
    print(f"{out.name}: {len(pdf)} pages → {kinds['p']} paragraphs, {kinds['Heading 1']} part headings, {kinds['Heading 2']} headings"
          f" (body {body}pt, {'Hindi' if hindi else 'English'})")


if __name__ == "__main__":
    for name in sys.argv[1:]:
        convert(name)
